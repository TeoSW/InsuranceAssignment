"""Read figures from a workbook: the sheet by its name, the column by its header, the unit from the
sheet header and the row by its label. No cell address is configured anywhere.
Every problem becomes a gap with a reason; nothing here stops the run."""

import openpyxl

from common import Value, gap, normalize_label, to_decimal


def read_figures(path, figures, settings, column="current"):
    """Each figure read from one column ("current" or "prior"). Returns ({id: Value}, hidden sheets)."""
    file = path.name
    try:
        book, formulas = open_book(path)
    except FileNotFoundError:
        return {figure["id"]: gap(figure, "file_missing", file=file) for figure in figures}, []
    except Exception as error:  # a damaged or foreign file must not stop the run
        return {figure["id"]: gap(figure, "file_unreadable", file=file, error=error) for figure in figures}, []
    hidden = [(sheet.title, sheet.sheet_state) for sheet in book.worksheets if sheet.sheet_state != "visible"]

    places = {figure["id"]: safely(figure, "căutarea", locate, book, figure, settings, column, file) for figure in figures}
    sharing = figures_sharing_a_row(places)
    values = {}
    for figure in figures:
        place = places[figure["id"]]
        if isinstance(place, Value):
            values[figure["id"]] = place
        elif figure["id"] in sharing:
            values[figure["id"]] = gap(figure, "row_two_figures", label=place["label_cell"].value,
                                       sheet=place["sheet"].title, file=file, figures=", ".join(sharing[figure["id"]]))
        else:
            values[figure["id"]] = safely(figure, "citirea", read_cell, place, formulas, figure, file)
    return values, hidden


def open_book(path):
    """The workbook twice: with the values as last saved, and with the formulas."""
    return openpyxl.load_workbook(path, data_only=True), openpyxl.load_workbook(path)


def read_elsewhere(path, figures, settings):
    """{figure id: [Value]} for the same labels found in the other visible sheets (this year's column).
    Used only to cross-check: a sheet where the label is absent or unreadable is skipped."""
    try:
        book, formulas = open_book(path)
    except Exception:
        return {}  # read_figures already reports the file
    found = {}
    for figure in figures:
        own = {normalize_label(name) for name in figure["sheets"]}
        for sheet in book.worksheets:
            if sheet.sheet_state != "visible" or normalize_label(sheet.title) in own:
                continue
            place = safely(figure, "căutarea", locate, book, dict(figure, sheets=[sheet.title]), settings, "current", path.name)
            if not isinstance(place, Value):
                value = safely(figure, "citirea", read_cell, place, formulas, figure, path.name)
                if value.status == "ok":
                    found.setdefault(figure["id"], []).append(value)
    return found


def figures_sharing_a_row(places):
    """{figure id: ids of all figures on its row}, for rows that more than one figure matched."""
    by_row = {}
    for figure_id, place in places.items():
        if not isinstance(place, Value):
            by_row.setdefault((place["sheet"].title, place["row"]), []).append(figure_id)
    return {figure_id: ids for ids in by_row.values() if len(ids) > 1 for figure_id in ids}


def safely(figure, step, function, *args):
    """Run one step for one figure; an unexpected error becomes a gap for that figure only."""
    try:
        return function(*args)
    except Exception as error:
        return gap(figure, "unexpected_error", what=f"{step} cifrei „{figure['id']}”", error=error)


def locate(book, figure, settings, column, file):
    """Where the figure is (sheet, row, column, label cell, sign, unit), or a gap saying why not."""
    label = figure["labels"][0]
    names = {normalize_label(name) for name in figure["sheets"]}
    sheets = [sheet for sheet in book.worksheets if normalize_label(sheet.title) in names]
    visible = [sheet for sheet in sheets if sheet.sheet_state == "visible"]
    if len(visible) != 1:
        reason = "sheet_ambiguous" if visible else "sheet_hidden" if sheets else "sheet_not_found"
        return gap(figure, reason, label=label, file=file, sheets=", ".join(figure["sheets"]))
    sheet = visible[0]
    here = {"file": file, "sheet": sheet.title}

    header = settings["reading"][column + "_column"]
    headers = text_cells(sheet, lambda text: normalize_label(header) in text)
    if len(headers) != 1:
        return gap(figure, "header_ambiguous" if headers else "header_not_found", label=label, header=header, **here)
    header_row, value_column = headers[0].row, headers[0].column

    rows = labelled_rows(sheet, figure, header_row)
    if not rows:
        return gap(figure, "label_not_found", label=label, **here)
    if len(rows) > 1:
        texts = [cell.value for cell, sign in rows]
        if len({normalize_label(text) for text in texts}) == 1:
            return gap(figure, "label_repeated", label=texts[0], count=len(rows), **here)
        return gap(figure, "alias_two_rows", label=label, rows=", ".join(f"„{text}”" for text in texts), **here)
    label_cell, sign = rows[0]

    place = {"sheet": sheet, "row": label_cell.row, "column": value_column, "label_cell": label_cell,
             "sign": sign, "unit": "", "lei": 1}
    if figure["kind"] == "amount":  # counts and percents never take the sheet unit
        unit, lei = sheet_unit(sheet, settings, header_row)
        if lei is None:
            reason = "unit_not_found" if unit is None else "unit_unknown"
            return gap(figure, reason, label=label_cell.value, unit=unit, **here)
        place["unit"], place["lei"] = unit, lei
    return place


def text_cells(sheet, test):
    """Cells holding text whose normalized form passes the test."""
    return [cell for row in sheet.iter_rows() for cell in row
            if isinstance(cell.value, str) and test(normalize_label(cell.value))]


def labelled_rows(sheet, figure, header_row):
    """(label cell, sign) for each row below the header that carries one of the figure's labels."""
    signs = {normalize_label(text): 1 for text in figure["labels"]}
    signs.update({normalize_label(text): -1 for text in figure.get("labels_stored_negative", [])})
    rows = {}
    for cell in text_cells(sheet, lambda text: text in signs):
        if cell.row > header_row and cell.row not in rows:
            rows[cell.row] = (cell, signs[normalize_label(cell.value)])
    return list(rows.values())


def sheet_unit(sheet, settings, header_row):
    """(unit text, lei per unit) from the "Valori în ..." cell above the header.
    The unit text is None if there is no such cell; lei is None if the unit is not in settings."""
    prefix = normalize_label(settings["reading"]["unit_prefix"])
    found = [normalize_label(cell.value)[len(prefix):].strip()
             for cell in text_cells(sheet, lambda text: text.startswith(prefix)) if cell.row <= header_row]
    if not found:
        return None, None
    unit = " / ".join(found)  # two unit cells give an unknown unit
    units = {normalize_label(name): lei for name, lei in settings["units"].items()}
    return unit, units.get(unit)


def read_cell(place, formulas, figure, file):
    """The figure's value from its cell, or a gap saying why the cell cannot be used."""
    sheet, row, column = place["sheet"], place["row"], place["column"]
    cell = sheet.cell(row, column)
    is_formula = formulas[sheet.title].cell(row, column).data_type == "f"
    source = f"{file}, foaia {sheet.title}, celula {cell.coordinate}" + (" (formulă)" if is_formula else "")
    label = place["label_cell"].value
    here = {"label": label, "file": file, "sheet": sheet.title}

    number = to_decimal(cell.value)
    if number is not None:
        return Value(figure["id"], figure["kind"], number * place["sign"] * place["lei"], source, label, place["unit"],
                     sheet=sheet.title)
    if is_formula:
        value = gap(figure, "formula_no_value", source, **here)
        value.formula = True
        return value
    if cell.value is not None and str(cell.value).strip():
        return gap(figure, "value_text", source, text=str(cell.value).strip(), **here)
    notes = [other.value.strip() for other in sheet[row]
             if isinstance(other.value, str) and other.value.strip() and other.column != place["label_cell"].column]
    if notes:
        return gap(figure, "value_missing_note", source, note="; ".join(notes), **here)
    return gap(figure, "value_missing", source, **here)
