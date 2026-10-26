"""Read the sources: figures from a workbook and quotes from documents.

A figure is found by the name of its sheet, the header of its column ("an curent" / "an anterior"),
the unit written above the header ("Valori în mii lei") and its row label. No cell address is
configured. Every problem becomes a gap with a reason; nothing here stops the run.
"""

import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from pypdf import PdfReader

# Old Romanian letters with a cedilla, mapped to the correct comma-below letters.
CEDILLA_TO_COMMA = str.maketrans({"ş": "ș", "Ş": "Ș", "ţ": "ț", "Ţ": "Ț"})
WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class Value:
    """One figure: its exact value (None for a gap), where it was found, and why it is a gap."""
    id: str
    kind: str                       # amount (in lei) | percent (a fraction) | count
    value: Decimal | None = None
    source: str = ""                # file, sheet and cell; document and page; or the calculation
    label: str = ""                 # the row label as written in the workbook
    unit: str = ""                  # the unit read from the sheet, for amounts
    status: str = "ok"              # ok | gap | conflict | decision
    reason: tuple = ()              # (reason name in gap_reasons.toml, {placeholder: text})
    candidates: list = field(default_factory=list)  # [(exact value, source)] when sources disagree
    why: list = field(default_factory=list)         # why this is the right value, for trasabilitate.csv


def gap(figure, reason, source="", **args):
    """A figure that cannot be filled, with a reason from gap_reasons.toml and its placeholders."""
    return Value(figure["id"], figure["kind"], None, source, args.get("label", ""), status="gap", reason=(reason, args))


def normalize(text):
    """Comma-below ș and ț, one space between words, no spaces at the ends."""
    return re.sub(r"\s+", " ", text.translate(CEDILLA_TO_COMMA)).strip()


def label_key(text):
    """How labels, sheet names and people are compared: normalized, ignoring case."""
    return normalize(text).casefold()


def to_decimal(cell_value):
    """A number from a workbook as an exact Decimal; None if the cell does not hold a number.
    A float keeps 15 significant digits, as Excel does: this drops the binary noise Excel saves
    (293699.39999999997 becomes 293699.4) and keeps every digit a person can type."""
    if isinstance(cell_value, bool) or not isinstance(cell_value, (int, float)):
        return None
    number = Decimal(format(cell_value, ".15g") if isinstance(cell_value, float) else cell_value)
    return number if number.is_finite() else None


def read_workbook(path, figures, settings, column="current"):
    """Each figure from one column ("current" or "prior"). Returns ({id: Value}, hidden sheets)."""
    try:
        book = openpyxl.load_workbook(path, data_only=True)  # the values as last saved
        formulas = openpyxl.load_workbook(path)  # the same cells, showing their formulas
    except Exception as error:  # a missing, damaged or foreign file must not stop the run
        reason = "file_missing" if isinstance(error, FileNotFoundError) else "file_unreadable"
        return {figure["id"]: gap(figure, reason, file=path.name, error=error) for figure in figures}, []
    hidden = [(sheet.title, sheet.sheet_state) for sheet in book.worksheets if sheet.sheet_state != "visible"]
    values = {}
    for figure in figures:
        try:
            values[figure["id"]] = read_figure(book, formulas, figure, settings, column, path.name)
        except Exception as error:  # a surprise on one figure is a gap for that figure only
            values[figure["id"]] = gap(figure, "unexpected_error", what=figure["id"], error=error)
    return values, hidden


def read_figure(book, formulas, figure, settings, column, file):
    """One figure: its visible sheet, its column, its row and its unit; or a gap saying which is missing."""
    label = figure["labels"][0]
    names = {label_key(name) for name in figure["sheets"]}
    sheets = [sheet for sheet in book.worksheets if label_key(sheet.title) in names and sheet.sheet_state == "visible"]
    if len(sheets) != 1:
        return gap(figure, "sheet_not_found", label=label, file=file, sheets=", ".join(figure["sheets"]))
    sheet = sheets[0]
    here = {"label": label, "file": file, "sheet": sheet.title}

    header = label_key(settings["reading"][column + "_column"])
    whole_phrase = re.compile(rf"(?<!\w){re.escape(header)}(?!\w)")  # "an curent", not "plan curent"
    headers = [cell for cell in text_cells(sheet) if whole_phrase.search(label_key(cell.value))]
    if len(headers) != 1:
        return gap(figure, "header_not_found", header=header, **here)
    header_row, value_column = headers[0].row, headers[0].column

    signs = {label_key(text): 1 for text in figure["labels"]}
    signs.update({label_key(text): -1 for text in figure.get("labels_stored_negative", [])})
    rows = [cell for cell in text_cells(sheet) if cell.row > header_row and label_key(cell.value) in signs]
    if len(rows) != 1:
        return gap(figure, "label_not_unique" if rows else "label_not_found", count=len(rows), **here)
    label_cell = rows[0]
    here["label"] = label_cell.value

    unit, lei = "", 1  # counts and percents never take the sheet unit
    if figure["kind"] == "amount":
        unit = sheet_unit(sheet, settings["reading"]["unit_prefix"], header_row)
        lei = {label_key(name): number for name, number in settings["units"].items()}.get(unit)
        if lei is None:
            return gap(figure, "unit_unknown", unit=unit, **here)

    cell = sheet.cell(label_cell.row, value_column)
    is_formula = formulas[sheet.title].cell(label_cell.row, value_column).data_type == "f"
    source = f"{file}, foaia {sheet.title}, celula {cell.coordinate}" + (" (formulă)" if is_formula else "")
    number = to_decimal(cell.value)
    if number is not None:
        why = [f"eticheta „{label_cell.value}” apare o singură dată în foaia vizibilă {sheet.title}",
               f"coloana „{headers[0].value}”"] + ([f"unitatea „{unit}” din antetul foii"] if unit else [])
        return Value(figure["id"], figure["kind"], number * signs[label_key(label_cell.value)] * lei, source,
                     label_cell.value, unit, why=why)
    if is_formula:
        return gap(figure, "formula_no_value", source, **here)
    if cell.value is not None and str(cell.value).strip():
        return gap(figure, "value_text", source, text=str(cell.value).strip(), **here)
    notes = [other.value.strip() for other in sheet[label_cell.row]
             if isinstance(other.value, str) and other.value.strip() and other.column != label_cell.column]
    if notes:
        return gap(figure, "value_missing_note", source, note="; ".join(notes), **here)
    return gap(figure, "value_missing", source, **here)


def text_cells(sheet):
    """Every cell of the sheet that holds text."""
    return [cell for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str)]


def sheet_unit(sheet, prefix, header_row):
    """The unit written after "Valori în" above the header ("mii lei"). Two such cells give an unknown unit."""
    prefix = label_key(prefix)
    found = [label_key(cell.value)[len(prefix):].strip() for cell in text_cells(sheet)
             if cell.row <= header_row and label_key(cell.value).startswith(prefix)]
    return " / ".join(found)


def document_text(path, page):
    """The normalized text of one page of a PDF, or of a whole .docx or .txt (they have no pages)."""
    if path.suffix == ".pdf":
        pages = PdfReader(path).pages
        return normalize(pages[page - 1].extract_text() or "") if 1 <= page <= len(pages) else ""
    if path.suffix == ".docx":
        root = ET.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
        return normalize(" ".join("".join(t.text or "" for t in p.iter(WORD + "t")) for p in root.iter(WORD + "p")))
    return normalize(path.read_text(encoding="utf-8"))


def check_fact(fact, folder):
    """None when the fact can be used: confirmed by a person, and its quote found on its page.
    Otherwise the gap reason. A fact without a document rests on the confirmation alone."""
    if not fact["confirmed_by"].strip():
        return ("fact_unconfirmed", {"statement": fact["statement"]})
    if not fact["document"]:
        return None
    try:
        text = document_text(folder / fact["document"], fact["page"])
    except Exception as error:
        return ("file_unreadable", {"file": fact["document"], "error": error})
    if normalize(fact["quote"]) not in text:
        return ("fact_not_found", {"statement": fact["statement"], "document": fact["document"], "page": fact["page"]})
    return None


def ro_number(text):
    """A number written the Romanian way ("15.000.000", "1.234,5") as a Decimal; None if it is not one."""
    try:
        number = Decimal(text.replace(".", "").replace(",", "."))
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def fact_figure(fact, problem, settings):
    """The figure a fact supplies (its value, written as in the quote), or a gap if the fact cannot be used."""
    figure = {"id": fact["figure"], "kind": "amount"}
    if problem:
        return gap(figure, problem[0], label=fact["statement"], **problem[1])
    lei = {label_key(name): number for name, number in settings["units"].items()}[label_key(fact["unit"])]
    number = ro_number(fact["value"]) * lei
    source = f"{fact['document']}, pagina {fact['page']}: „{fact['quote']}”"
    why = [f"citatul a fost găsit în {fact['document']}, pagina {fact['page']}",
           f"confirmat de {fact['confirmed_by']} la {fact['confirmed_on']}"]
    return Value(fact["figure"], "amount", number, source, fact["statement"], fact["unit"], why=why)
