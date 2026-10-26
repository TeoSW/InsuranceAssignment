"""Mutation tests (T10): seven changes to a copy of registru_2025.xlsx, each run through the whole program
and compared with the run on the original. A wrong figure must never reach the text; a real problem must
become a gap or a conflict; a harmless change must give the same draft.
Print the table:   python -X utf8 tests/test_mutations.py
"""

import csv
import re
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest import mock

import openpyxl

import helpers
import main

GAP = re.compile(r"\[de completat:[^\]]*\]")
NUMBER = re.compile(r"\d+(?:[.,]\d+)*")


def cell_with(sheet, text):
    return next(cell for row in sheet.iter_rows() for cell in row if cell.value == text)


def moved_rows(book):
    sheet = book["Capital"]
    rows = [list(row) for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=2, max_col=5, values_only=True)]
    for offset, values in enumerate(reversed(rows)):
        for column, value in enumerate(values, 2):
            sheet.cell(4 + offset, column).value = value  # cell(..., value) would not clear an empty cell


def hide_capital(book):
    book["Capital"].sheet_state = "hidden"


def known_alias(book):
    label = cell_with(book["SCR_module"], "Efect de diversificare (se scade)")
    label.value = "Diversificare (valoare negativă)"
    value = book["SCR_module"].cell(label.row, label.column + 1)
    value.value = -value.value


def unknown_label(book):
    cell_with(book["Capital"], "SCR").value = "Cerința de capital de solvabilitate (SCR)"


def unit_scaled(book):
    sheet = book["Capital"]
    cell_with(sheet, "Valori în mii lei").value = "Valori în lei"
    for row in sheet.iter_rows(min_row=4):
        if row[1].value != "Rata de solvabilitate":  # a ratio has no unit
            for cell in row[2:4]:
                if isinstance(cell.value, (int, float)):
                    cell.value = int(Decimal(str(cell.value)) * 1000)


def stale_mcr(book):
    label = cell_with(book["Capital"], "MCR")
    book["Capital"].cell(label.row, label.column + 1).value = 45570  # the 2024 MCR, in mii lei


# name, what changes, change, keep formulas (save as a script would), figures that must be caught, new gaps allowed
MUTATIONS = [
    ("rânduri mutate", "rândurile din Capital în ordine inversă", moved_rows, False, [], 0),
    ("foaie ascunsă", "foaia Capital ascunsă", hide_capital, False, ["own_funds_total", "scr"], 2),
    ("etichetă: alt alias", "diversificarea cu eticheta și semnul din 2024", known_alias, False, [], 0),
    ("etichetă necunoscută", "„SCR” din Capital redenumit", unknown_label, False, ["scr"], 1),
    ("unitate schimbată", "Capital din mii lei în lei, valorile × 1.000", unit_scaled, False, [], 0),
    ("formulă fără valoare", "registrul salvat de un script, fără valori salvate", None, True, [], 0),
    ("cifră din 2024", "MCR completat cu cifra din 2024 (45.570 mii lei)", stale_mcr, False, ["mcr"], 0),
]


def run_project(folder, change=None, keep_formulas=False):
    """The whole program on a copy of the project (facts confirmed, as at checkpoint 2). Returns (draft, statuses)."""
    helpers.project_copy(folder, confirm=True)
    if change or keep_formulas:
        path = folder / "date" / "registru_2025.xlsx"
        book = openpyxl.load_workbook(path, data_only=not keep_formulas)
        if change:
            change(book)
        book.save(path)
    with mock.patch("builtins.print"):
        main.main(root=folder, folder=folder / "date")
    draft = (folder / "output" / "sectiunea_E_2025.md").read_text(encoding="utf-8")
    with open(folder / "output" / "trasabilitate.csv", encoding="utf-8-sig", newline="") as file:
        statuses = {row["cifra"]: row["stare"] for row in csv.DictReader(file, delimiter=";")}
    return draft, statuses


def wrong_figures(draft, original):
    """Numbers in text outside gaps that are not in the same line of the original draft."""
    wrong = []
    for line, original_line in zip(draft.splitlines(), original.splitlines(), strict=True):
        for fragment in GAP.split(line):
            if fragment.strip() and fragment.strip() not in original_line:
                wrong += NUMBER.findall(fragment)
    return wrong


def run_all():
    results = []
    with tempfile.TemporaryDirectory() as tmp:
        original, _ = run_project(Path(tmp) / "original")
        for number, (name, what, change, keep, must_catch, gaps_allowed) in enumerate(MUTATIONS, 1):
            draft, statuses = run_project(Path(tmp) / str(number), change, keep)
            new_gaps = draft.count("[de completat") - original.count("[de completat")
            results.append({"name": name, "what": what, "wrong": wrong_figures(draft, original),
                            "caught": all(statuses[i] != "ok" for i in must_catch) if must_catch else None,
                            "expected": "gol sau conflict" if must_catch else "aceeași cifră",
                            "new_gaps": new_gaps, "gaps_allowed": gaps_allowed,
                            "unnecessary": max(0, new_gaps - gaps_allowed)})
    return results


def table(results):
    lines = ["| Mutație | Ce s-a schimbat | Așteptat | Cifre greșite în text | Detectat | Goluri noi (permise) | Goluri inutile |",
             "|---|---|---|---|---|---|---|"]
    for r in results:
        caught = "–" if r["caught"] is None else "da" if r["caught"] else "NU"
        lines.append(f"| {r['name']} | {r['what']} | {r['expected']} | {len(r['wrong'])} | {caught} | "
                     f"{r['new_gaps']} ({r['gaps_allowed']}) | {r['unnecessary']} |")
    problems = [r for r in results if r["caught"] is not None]
    lines += ["", f"Cifre greșite în text: {sum(len(r['wrong']) for r in results)}. "
                  f"Rata de detecție: {sum(r['caught'] for r in problems)}/{len(problems)}. "
                  f"Goluri inutile: {sum(r['unnecessary'] for r in results)}."]
    return "\n".join(lines)


class WrongFigureMetricTest(unittest.TestCase):
    def test_the_metric_sees_a_wrong_number_outside_gaps_only(self):
        self.assertEqual(wrong_figures("SCR 151,9 mil. lei. [de completat: 45,6]", "SCR 148,8 mil. lei. [x]"), ["151,9"])


@unittest.skipUnless((helpers.DATE / "registru_2025.xlsx").exists(), "date/ lipsește")
class MutationTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.results = run_all()

    def test_no_wrong_figure_reaches_the_text(self):
        self.assertEqual({r["name"]: r["wrong"] for r in self.results if r["wrong"]}, {})

    def test_every_real_problem_is_caught(self):
        self.assertEqual([r["name"] for r in self.results if r["caught"] is False], [])

    def test_no_unnecessary_gap(self):
        self.assertEqual({r["name"]: r["new_gaps"] for r in self.results if r["new_gaps"] != r["gaps_allowed"]}, {})


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    print(table(run_all()))
